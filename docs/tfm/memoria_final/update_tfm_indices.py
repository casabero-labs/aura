from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document


def normalise(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def locate_page(target: str, pages: list[str]) -> int | None:
    needle = normalise(target)
    candidates = [
        index
        for index, page in enumerate(pages, start=1)
        if index > 6 and needle in normalise(page)
    ]
    if candidates:
        return max(candidates)
    short = needle[:55]
    candidates = [
        index
        for index, page in enumerate(pages, start=1)
        if index > 6 and short in normalise(page)
    ]
    return max(candidates) if candidates else None


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: update_tfm_indices.py document.docx rendered.txt")
    docx_path = Path(sys.argv[1])
    text_path = Path(sys.argv[2])
    pages = text_path.read_text(encoding="utf-8", errors="ignore").split("\f")
    doc = Document(docx_path)

    start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "Índice de contenidos")
    end = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "1. Introducción")
    unresolved: list[str] = []
    updated = 0
    for paragraph in doc.paragraphs[start + 1:end]:
        if "\t" not in paragraph.text:
            continue
        target = paragraph.text.rsplit("\t", 1)[0].strip()
        page = locate_page(target, pages)
        if page is None:
            unresolved.append(target)
            continue
        if paragraph.runs:
            paragraph.runs[0].text = f"{target}\t{page}"
            for extra in paragraph.runs[1:]:
                extra.text = ""
        else:
            paragraph.text = f"{target}\t{page}"
        updated += 1

    doc.save(docx_path)
    print(f"updated={updated}")
    if unresolved:
        print("unresolved:")
        for value in unresolved:
            print(f"- {value}")
        raise SystemExit(2)


if __name__ == "__main__":
    main()
