from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.shared import Inches

REPO = Path(__file__).resolve().parents[2]
DOCX = REPO / "docs/memoria/entregas/segunda_entrega/Segunda_Entrega_TFM_Joseph_Gari_Borrador_Estructurado.docx"
RESULTS = REPO / "docs/evidence/results/aura_evidence_results.json"
FIGURES = REPO / "docs/evidence/screenshots/docx_figures"

FIGURE_MAP = {
    "Figura 5.2.": ("aura-clientes-perfil.png", "Figura 5.2. Perfilamiento actualizado del dataset clientes_sucio.csv: caracterización, score y hallazgos reproducibles bajo la interfaz Casabero."),
    "Figura 5.3.": ("aura-inventario-perfil.png", "Figura 5.3. Perfilamiento actualizado del dataset inventario_sucio.csv con hallazgos de integridad, tipos y consistencia."),
    "Figura 5.4.": ("aura-operaciones-perfil.png", "Figura 5.4. Perfilamiento actualizado del dataset operaciones_sucio.csv con duplicados, nulos y reglas operacionales."),
    "Figura 5.5.": ("aura-clientes-diagnostico.png", "Figura 5.5. Etapa de diagnóstico asistido: contrato de entrada, problema observado y paquete controlado para el LLM."),
    "Figura 5.6.": ("aura-clientes-script.png", "Figura 5.6. Laboratorio experimental de modelos: protocolo, dataset, contrato AURA, runner y matriz de benchmark."),
    "Figura 5.7.": ("aura-clientes-revision.png", "Figura 5.7. Drawer de configuración: selección de proveedor local/cloud y edición guiada del contrato técnico."),
}


def remove_picture_paragraphs(doc: Document) -> None:
    for paragraph in list(doc.paragraphs):
        if paragraph._p.xpath(".//w:drawing") or paragraph._p.xpath(".//w:pict"):
            paragraph._element.getparent().remove(paragraph._element)


def update_evidence_table(doc: Document, rows: list[dict]) -> None:
    if len(doc.tables) < 2:
        return
    table = doc.tables[1]
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for row in rows:
        cells = table.add_row().cells
        cells[0].text = row["dataset"]
        cells[1].text = str(row["rows"])
        cells[2].text = str(row["columns"])
        cells[3].text = str(row["issues"])
        cells[4].text = f'{row["score"]}/100'


def replace_text(doc: Document) -> None:
    replacements = {
        "Para esta entrega se ejecutó AURA con tres datasets pequeños incluidos en el repositorio: clientes_sucio.csv, inventario_sucio.csv y operaciones_sucio.csv. Estos datasets activan reglas de duplicados, placeholders, espacios y formatos inválidos.":
            "Para esta entrega se ejecutó nuevamente AURA con tres datasets pequeños incluidos en el repositorio: clientes_sucio.csv, inventario_sucio.csv y operaciones_sucio.csv. Las capturas fueron regeneradas después de la migración visual al estándar Casabero, por lo que muestran la interfaz actual de carga, perfilamiento, diagnóstico controlado, laboratorio experimental y configuración del contrato técnico.",
        "Figura 5.6. Generación de script de asistencia con respaldo determinista y matriz de validación automática.":
            "Figura 5.6. Laboratorio experimental de modelos: protocolo, dataset, contrato AURA, runner y matriz de benchmark.",
        "Figura 5.7. Revisión humana del script, decisión de aprobación y preparación para simulación.":
            "Figura 5.7. Drawer de configuración: selección de proveedor local/cloud y edición guiada del contrato técnico.",
        "La evaluación preliminar se centra en dos dimensiones: aplicabilidad técnica y usabilidad del flujo.":
            "La evaluación preliminar se centra en tres dimensiones: aplicabilidad técnica, usabilidad del flujo y alineación visual-documental tras la migración de interfaz.",
        "El flujo revisado reduce ruido visual y evita que el usuario tenga que interpretar nombres internos de capas o prompts extensos.":
            "El flujo revisado reduce ruido visual, agrupa la información por responsabilidad y evita que el usuario tenga que interpretar nombres internos de capas o prompts extensos.",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text
        for old, new in replacements.items():
            if old in text:
                paragraph.text = text.replace(old, new)
                text = paragraph.text
        for prefix, (_, caption) in FIGURE_MAP.items():
            if paragraph.text.startswith(prefix):
                paragraph.text = caption


def insert_figures(doc: Document) -> None:
    for paragraph in list(doc.paragraphs):
        for prefix, (filename, caption) in FIGURE_MAP.items():
            if paragraph.text.startswith(prefix):
                paragraph.text = caption
                run = doc.add_paragraph().add_run()
                run.add_picture(str(FIGURES / filename), width=Inches(6.4))
                paragraph._p.addnext(run._parent._p)
                break


def main() -> None:
    payload = json.loads(RESULTS.read_text())
    doc = Document(DOCX)
    remove_picture_paragraphs(doc)
    replace_text(doc)
    update_evidence_table(doc, payload["rows"])
    insert_figures(doc)
    doc.save(DOCX)
    print(f"Documento actualizado: {DOCX}")


if __name__ == "__main__":
    main()
